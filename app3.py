"""
Style-Locked AI Slide Generator (v2)
====================================
Topic -> AI plans slides -> Art Director reads your style image and writes a
style spec -> an image model paints EACH WHOLE SLIDE in that style.

Styles : put images in  styles/  next to this file (styles/skillshift.png).
         Optional styles/skillshift.txt = saved style spec (auto-created text
         can be pasted there so it never has to be regenerated).
Keys   : Gemini key (FREE, https://aistudio.google.com/apikey) is needed for
         planning + art director. Painting can use Gemini (free), OpenAI
         (paid) or Pollinations (free key from enter.pollinations.ai).
Install: pip install streamlit google-genai pillow python-pptx requests openai
"""
import base64
import io
import json
import re
import zipfile
from pathlib import Path
from urllib.parse import quote

import requests
import streamlit as st
from google import genai
from google.genai import types
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

try:
    from openai import OpenAI
except ImportError:
    OpenAI = None

STYLES_DIR = Path(__file__).parent / "styles"
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp"}
PROVIDERS = ["Gemini (free tier)", "Pollinations (free key)", "OpenAI gpt-image-1 (paid)"]

PLANNER = """You are a presentation designer. Return ONLY raw JSON (no fences):
{"deck_title": "...", "slides": [
  {"role": "cover|content|summary",
   "headline": "short, max 6 words",
   "points": ["2-4 very short phrases, max 8 words each"],
   "visual": "what the main illustration/diagram/infographic should show"}
]}
Exactly N slides. Slide 1 role=cover (points = one tagline). Last slide role=summary.
All text SHORT (it is rendered inside an image). Plain text only."""

DIRECTOR = """You are an art director. Study this reference slide and write a precise STYLE SPECIFICATION
(max 350 words) that another image model can follow to reproduce the same look on a different topic.
Cover: 1) canvas + background atmosphere, 2) color palette with hex codes, 3) typography (title weight/size,
two-tone or highlighted words, letter-spacing, sub-heading style), 4) layout zones and grid (header, numbered
cards, icon rows, banners, side panels, footer/corners) with rough positions, 5) illustration/icon rendering
style, 6) recurring UI elements (numbered badges, dividers, arrows, frames), 7) text density.
Describe style and structure ONLY. Never mention the topic or literal words of the reference."""


# ------------------------------------------------------------ styles
def load_styles():
    out = {}
    if STYLES_DIR.exists():
        for p in sorted(STYLES_DIR.iterdir()):
            if p.suffix.lower() in IMG_EXT:
                t = p.with_suffix(".txt")
                out[p.stem] = {"image": Image.open(p).convert("RGB"),
                               "spec": t.read_text(encoding="utf-8") if t.exists() else ""}
    return out


def art_director(client, model, img):
    return client.models.generate_content(model=model, contents=[DIRECTOR, img]).text.strip()


# ------------------------------------------------------------ planning
def plan_deck(client, model, topic, n, lang):
    r = client.models.generate_content(model=model, contents=f"{PLANNER}\n\nN = {n}\nLanguage: {lang}\nTopic:\n{topic}")
    raw = re.sub(r"^```(json)?|```$", "", r.text.strip(), flags=re.I | re.M).strip()
    plan = json.loads(raw[raw.find("{"): raw.rfind("}") + 1])
    if not plan.get("slides"):
        raise ValueError("Empty plan, try again.")
    return plan


# ------------------------------------------------------------ prompt
ROLE_HINT = {
    "cover": "Big bold title as the hero (use the reference's title treatment), tagline under it, "
             "a row of small circular illustrations related to the topic.",
    "content": "Use the reference's card system: numbered header cards / icon rows for the points, "
               "plus a side or bottom illustrated example panel if the reference has one.",
    "summary": "Finish with the reference's bottom callout banner carrying the key takeaway, "
               "plus a compact recap of the points.",
}


def slide_prompt(slide, deck_title, idx, total, spec, footer):
    pts = "\n".join(f"- {p}" for p in slide.get("points", []))
    return f"""Create ONE 16:9 presentation slide as a finished infographic-style image.

STYLE REFERENCE: match the attached reference image as closely as possible: same palette, typography,
panel/card structure, icon and illustration rendering, background atmosphere, spacing and density.
Do NOT copy its words, logos, people or subject.
STYLE SPEC (authoritative):
{spec}

LAYOUT: {ROLE_HINT.get(slide.get('role', 'content'), ROLE_HINT['content'])}

TEXT TO RENDER (spelled exactly, legible, nothing else):
Headline: {slide.get('headline', '')}
{pts}
Small label: {deck_title} | {idx + 1}/{total}{f' | {footer}' if footer else ''}
MAIN VISUAL: {slide.get('visual', '')}
No watermarks. Keep everything inside the frame."""


# ------------------------------------------------------------ painters
def paint_gemini(key, model, style_img, prev, prompt):
    c = genai.Client(api_key=key)
    contents = [prompt, style_img]
    if prev is not None:
        contents += ["Also an earlier slide of THIS deck, stay consistent with it:", prev]
    try:
        cfg = types.GenerateContentConfig(response_modalities=["IMAGE"],
                                          image_config=types.ImageConfig(aspect_ratio="16:9"))
    except Exception:
        cfg = types.GenerateContentConfig(response_modalities=["IMAGE"])
    r = c.models.generate_content(model=model, contents=contents, config=cfg)
    for cand in r.candidates or []:
        for part in cand.content.parts or []:
            if getattr(part, "inline_data", None) and part.inline_data.data:
                return Image.open(io.BytesIO(part.inline_data.data)).convert("RGB")
    raise RuntimeError("Gemini returned no image (quota, safety block, or model not on your tier).")


def _buf(img, name):
    b = io.BytesIO()
    img.save(b, "PNG")
    b.name = name
    b.seek(0)
    return b


def paint_openai(key, model, style_img, prev, prompt):
    if OpenAI is None:
        raise RuntimeError("pip install openai")
    files = [_buf(style_img, "style.png")] + ([_buf(prev, "prev.png")] if prev is not None else [])
    r = OpenAI(api_key=key).images.edit(model=model, image=files, prompt=prompt, size="1536x1024")
    return Image.open(io.BytesIO(base64.b64decode(r.data[0].b64_json))).convert("RGB")


def paint_pollinations(key, model, style_img, prev, prompt):
    # Style spec carries the style (reference upload needs a public URL, so it is not sent).
    url = f"https://gen.pollinations.ai/image/{quote(prompt[:1800])}?model={model}&width=1344&height=768&key={key}"
    r = requests.get(url, timeout=180)
    r.raise_for_status()
    return Image.open(io.BytesIO(r.content)).convert("RGB")


def paint(provider, keys, models, style_img, prev, prompt):
    if provider == PROVIDERS[0]:
        return paint_gemini(keys["gemini"], models["gemini"], style_img, prev, prompt)
    if provider == PROVIDERS[1]:
        return paint_pollinations(keys["poll"], models["poll"], style_img, prev, prompt)
    return paint_openai(keys["openai"], models["openai"], style_img, prev, prompt)


# ------------------------------------------------------------ export
def png(img):
    b = io.BytesIO()
    img.save(b, "PNG")
    return b.getvalue()


def build_pptx(imgs):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for im in imgs:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(io.BytesIO(png(im)), 0, 0, prs.slide_width, prs.slide_height)
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def build_pdf(imgs):
    b = io.BytesIO()
    imgs[0].save(b, "PDF", save_all=True, append_images=imgs[1:], resolution=150)
    return b.getvalue()


def build_zip(imgs):
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        for i, im in enumerate(imgs, 1):
            z.writestr(f"slide_{i:02d}.png", png(im))
    return b.getvalue()


# ------------------------------------------------------------ app
def main():
    st.set_page_config(page_title="Style-Locked Slide Generator", layout="wide")
    st.title("🎨 Style-Locked AI Slide Generator")
    ss = st.session_state
    ss.setdefault("plan", None)
    ss.setdefault("images", [])
    ss.setdefault("specs", {})

    def secret(k):
        try:
            return st.secrets.get(k, "")
        except Exception:
            return ""

    with st.sidebar:
        st.header("Keys")
        gkey = st.text_input("Gemini key (free) - planning + art director", type="password", value=secret("GEMINI_API_KEY"))
        provider = st.selectbox("Image painter", PROVIDERS)
        keys, models = {"gemini": gkey, "poll": "", "openai": ""}, {
            "gemini": "gemini-2.5-flash-image", "poll": "nanobanana", "openai": "gpt-image-1"}
        if provider == PROVIDERS[1]:
            keys["poll"] = st.text_input("Pollinations key (free at enter.pollinations.ai)", type="password", value=secret("POLLINATIONS_KEY"))
            models["poll"] = st.text_input("Pollinations image model", models["poll"], help="See gen.pollinations.ai/image/models")
        elif provider == PROVIDERS[2]:
            keys["openai"] = st.text_input("OpenAI key (needs paid billing)", type="password", value=secret("OPENAI_API_KEY"))
        else:
            models["gemini"] = st.text_input("Gemini image model", models["gemini"],
                                             help="Try a newer image model if your tier allows it. Better models render text better.")
        text_model = st.text_input("Gemini text model", "gemini-2.5-flash")

        st.header("Style")
        styles = load_styles()
        up = st.file_uploader("Upload style image", type=["png", "jpg", "jpeg", "webp"])
        if up:
            styles = {"Uploaded style": {"image": Image.open(up).convert("RGB"), "spec": ""}, **styles}
        if not styles:
            st.warning("Add images to a `styles/` folder or upload one.")
            st.stop()
        name = st.selectbox("Pick style", list(styles))
        style = styles[name]
        st.image(style["image"], use_container_width=True)
        if name not in ss.specs:
            ss.specs[name] = style["spec"]
        if st.button("🎬 Analyze style with Art Director", use_container_width=True):
            if not gkey:
                st.warning("Gemini key needed.")
            else:
                with st.spinner("Reading the style..."):
                    try:
                        ss.specs[name] = art_director(genai.Client(api_key=gkey), text_model, style["image"])
                    except Exception as e:
                        st.error(e)
        ss.specs[name] = st.text_area("Style spec (edit freely)", ss.specs[name], height=220, key=f"spec_{name}")
        if not ss.specs[name]:
            st.info("Click Analyze - the spec is what makes results match your reference.")

        st.header("Deck")
        n = st.slider("Slides", 3, 12, 6)
        lang = st.text_input("Slide language", "English")
        footer = st.text_input("Footer text", "")
        chain = st.checkbox("Chain slides for consistency", True)

    spec = ss.specs[name]
    topic = st.text_area("Topic", height=110, placeholder="e.g. Tool vs task vs skill in architecture")

    if st.button("🧠 1. Plan slides", type="primary", use_container_width=True):
        if not gkey or not topic.strip():
            st.warning("Need the Gemini key and a topic.")
        else:
            with st.spinner("Planning..."):
                try:
                    ss.plan = plan_deck(genai.Client(api_key=gkey), text_model, topic, n, lang)
                    ss.images = []
                except Exception as e:
                    st.error(f"Planning failed: {e}")

    plan = ss.plan
    if not plan:
        return
    st.subheader("Review text (it is painted into the image)")
    plan["deck_title"] = st.text_input("Deck title", plan.get("deck_title", ""))
    for i, s in enumerate(plan["slides"]):
        with st.expander(f"Slide {i + 1}: {s.get('headline', '')}"):
            s["headline"] = st.text_input("Headline", s.get("headline", ""), key=f"h{i}")
            p = st.text_area("Points (one per line)", "\n".join(s.get("points", [])), key=f"p{i}")
            s["points"] = [x.strip() for x in p.splitlines() if x.strip()]
            s["visual"] = st.text_area("Main visual", s.get("visual", ""), key=f"v{i}", height=70)

    def one(i, prev):
        s, total = plan["slides"][i], len(plan["slides"])
        pr = slide_prompt(s, plan["deck_title"], i, total, spec, footer)
        return paint(provider, keys, models, style["image"], prev if chain else None, pr)

    if st.button("🖌️ 2. Paint all slides", type="primary", use_container_width=True):
        imgs, bar = [], st.progress(0.0)
        for i in range(len(plan["slides"])):
            bar.progress(i / len(plan["slides"]), f"Painting {i + 1}/{len(plan['slides'])}...")
            try:
                imgs.append(one(i, imgs[0] if imgs else None))
            except Exception as e:
                st.error(f"Slide {i + 1} failed: {e}")
                imgs.append(Image.new("RGB", (1920, 1080), (30, 30, 30)))
        bar.empty()
        ss.images = imgs

    imgs = ss.images
    if imgs:
        st.subheader("Result")
        for i, im in enumerate(imgs):
            st.image(im, caption=f"Slide {i + 1}", use_container_width=True)
            if st.button(f"🔁 Regenerate slide {i + 1}", key=f"r{i}"):
                with st.spinner("Repainting..."):
                    try:
                        imgs[i] = one(i, imgs[0] if i > 0 else None)
                        st.rerun()
                    except Exception as e:
                        st.error(e)
        c1, c2, c3 = st.columns(3)
        c1.download_button("⬇️ PowerPoint", build_pptx(imgs), "deck.pptx", use_container_width=True)
        c2.download_button("⬇️ PDF", build_pdf(imgs), "deck.pdf", use_container_width=True)
        c3.download_button("⬇️ PNG ZIP", build_zip(imgs), "slides.zip", use_container_width=True)


if __name__ == "__main__":
    main()
